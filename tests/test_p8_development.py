import json
import shutil
from pathlib import Path
from uuid import uuid4

import httpx
import pytest

from agent_ops.deepseek import DeepSeekConfig, DeepSeekEvaluation, EvaluationUnavailable
from agent_ops.ecd_redaction import redact_ecd
from agent_ops.routing_evaluation import DevelopmentSample, development_versions, request_messages


def test_redaction_occurrences_idempotence_and_cues_retained():
    phone='1'+'38'+'0'*8
    order='8'*18
    raw=f'请联系{phone}，邮箱test@example.invalid，订单号:{order}，收货地址:示例省示例市示例区示例路8号。'
    result,counts=redact_ecd(raw)
    assert phone not in result and order not in result and 'example.invalid' not in result
    assert '示例路' not in result
    assert '[MOBILE]' in result and '[EMAIL]' in result and '订单号:[ORDER]' in result
    assert counts['mobile']==counts['email']==counts['order_labeled']==counts['address_labeled']==1
    assert redact_ecd(result)==(result,{})


async def test_ecd_requires_both_approval_and_clean_scan_and_cache_is_shared():
    path=Path('.local/test-p8-development')/uuid4().hex
    count=0
    def send(_):
        nonlocal count
        count+=1
        return httpx.Response(200,json={'model':'fixture','usage':{'prompt_tokens':2,'completion_tokens':1},
          'choices':[{'message':{'content':'{"route":"clarify"}'}}]})
    try:
        async with httpx.AsyncClient(transport=httpx.MockTransport(send)) as client:
            service=DeepSeekEvaluation(DeepSeekConfig('fixture','https://provider.invalid/c','fixture'),path,client=client,allow_network=True)
            body={'dataset_digest':'dev','source':'ecd_redacted','messages':[{'role':'user','content':'邮箱test@example.invalid'}]}
            with pytest.raises(EvaluationUnavailable,match='NOT_AUTHORIZED'):
                await service.complete(**body)
            service.allow_ecd=True
            with pytest.raises(EvaluationUnavailable,match='REDACTION_REQUIRED'):
                await service.complete(**body)
            body['messages']=[{'role':'user','content':'邮箱[EMAIL]'}]
            first=await service.complete(**body)
            assert first==await service.complete(**body) and count==1
            assert 'fixture' not in json.dumps(first['input_summary'])
    finally:shutil.rmtree(path,ignore_errors=True)


def test_three_versions_never_need_extra_model_result_and_no_final_metrics():
    def sample(text):
        return DevelopmentSample(id='fixture',source='p2_synthetic',message=text,context=[],
                                 group_hash='fixture',speaker_role='synthetic_customer')
    easy=development_versions(sample('请查询物流'),'order_status')
    assert not easy['hybrid_model_used'] and easy['hybrid']=='order_status'
    denial=development_versions(sample('不要人工，只查物流'),'order_status')
    assert denial['hybrid_model_used'] and denial['hybrid']=='order_status'
    assert 'accuracy' not in denial and denial['status'].endswith('not_frozen')
    message=request_messages(sample('测试'),'classifier')
    assert 'context' in message[1]['content'] and 'label' not in json.loads(message[1]['content'])


def test_annotation_workbook_is_blind_and_has_instructions_first(tmp_path):
    import runpy

    from openpyxl import load_workbook
    module=runpy.run_path('scripts/prepare-p8-development.py',run_name='fixture_exporter')
    samples=[]
    for source,n in [('p2_synthetic',40),('ecd_redacted',60)]:
        for i in range(n):
            samples.append({'id':f'{source}-{i}','source':source,'speaker_role':'unknown',
                            'message':'=literal text','context':['fixture context']})
    path=tmp_path/'blind.xlsx'
    ids=module['write_blind_workbook'](path,samples)
    book=load_workbook(path)
    assert len(ids)==len(set(ids))==100
    assert book.sheetnames==['01_标注说明','02_盲标100条']
    table=book.worksheets[1]
    assert table.max_row==101 and table.max_column==8
    for row in table.iter_rows(min_row=2):
        assert all(c.value is None for c in row[5:])
        assert row[4].data_type=='s'  # Leading = is literal, never a spreadsheet formula.
    assert all(s.sheet_state=='visible' for s in book)
    assert all(not d.hidden for d in table.column_dimensions.values())



def test_json_transport_scan_does_not_cross_fields_but_checks_nested_identifiers():
    from agent_ops.ecd_redaction import outbound_is_redacted
    assert outbound_is_redacted([{'content':json.dumps({'message':'地址呢','context':[]},ensure_ascii=False)}])
    assert not outbound_is_redacted([{'content':json.dumps({'context':[{'text':'test@example.invalid'}]})}])
    assert not outbound_is_redacted([{'content':json.dumps({'number':int('8'*18)})}])


async def test_development_worker_limit_and_model_hybrid_share_classification(tmp_path,monkeypatch):
    import asyncio
    import hashlib
    import runpy
    from types import SimpleNamespace

    from agent_ops.p8_pilot import SILVER_PROMPT
    module=runpy.run_path('scripts/run-p8-development.py',run_name='fixture_runner')
    rows=[]
    for source,n in [('p2_synthetic',800),('ecd_redacted',1200)]:
        for i in range(n):
            rows.append({'id':f'{source}-{i}','source':source,'message':f'fixture-{source}-{i}',
                         'context':[],'group_hash':'fixture','speaker_role':'unknown'})
    path=tmp_path/'dataset.json';path.write_text(json.dumps({'samples':rows}))
    summaries={};active=peak=0
    class Fake:
        def __init__(self,config,*_,**kw):
            self.config=config;self.network_requests=self.retries=self.rate_limits=self.cache_hits=0
            assert kw['allow_ecd'] and kw['allow_network']
        async def complete(self,**kw):
            nonlocal active,peak
            active+=1;peak=max(peak,active)
            await asyncio.sleep(0)
            active-=1;self.network_requests+=1
            silver=kw['messages'][0]['content']==SILVER_PROMPT
            content={'label':'clarify','uncertain':True,'basis_code':'AMBIGUOUS'} if silver else {'route':'clarify','reason_code':'AMBIGUOUS'}
            return {'fingerprint':hashlib.sha256(json.dumps(kw['messages']).encode()).hexdigest(),
                    'returned_model':'fixture','called_at':'2026-01-01T00:00:00+00:00',
                    'content':json.dumps(content),'usage':{'prompt_tokens':10,'completion_tokens':5}}
        async def aclose(self):pass
    def capture(path,value):
        if path.name in {'summary.json','development-versions.json'}:summaries[path.name]=value
    run=module['run'];monkeypatch.setitem(run.__globals__,'DeepSeekEvaluation',Fake)
    monkeypatch.setitem(run.__globals__,'private_write',capture)
    for k,v in {'DEEPSEEK_API_KEY':'fixture','AGENT_OPS_DEEPSEEK_ENDPOINT':'https://provider.invalid/c',
                'AGENT_OPS_DEEPSEEK_MODEL':'fixture','AGENT_OPS_LLM_CONCURRENCY':'2'}.items():monkeypatch.setenv(k,v)
    await run(SimpleNamespace(dataset=path,output=tmp_path,cache=None,allow_network=True,allow_ecd=True))
    summary=summaries['summary.json']
    assert summary['jobs']==summary['network_requests']==3200 and peak<=2
    assert len(summaries['development-versions.json'])==2000
    assert summary['hybrid_additional_requests']==0 and not summary['final_test_sampled']
    assert all(g['failed_or_invalid']==0 for g in summary['groups'].values())
