import pytest

from agent_ops.p8_release import register_round, reserve_evaluation
from agent_ops.routing_metrics import wilson_interval
from agent_ops.routing_policy import evaluation_rule


@pytest.mark.parametrize('text',['请问用哪家快递','可以指定甲快递吗','给我改用乙物流','能发甲承运商吗','发什么快递'])
def test_carrier_options_not_existing_order_progress(text):
    assert evaluation_rule(text)=='clarify'


@pytest.mark.parametrize('text',['请查询快递到哪里了','已经发出的包裹何时送达','我要人工投诉快递','这单我要退款'])
def test_existing_business_actions_preserved(text):
    assert evaluation_rule(text) in {'order_status','human','return_request'}


def test_wilson_denominators_and_boundaries():
    assert wilson_interval(0,0)['low'] is None
    assert wilson_interval(0,10)['high']==pytest.approx(.2775328,abs=1e-6)
    assert wilson_interval(10,10)['low']==pytest.approx(.7224672,abs=1e-6)
    assert wilson_interval(50,100)['low']==pytest.approx(.4038315,abs=1e-6)


def test_three_round_limit_and_once_guard(tmp_path):
    path=tmp_path/'rounds.json'
    register_round(path,1,'first');register_round(path,1,'first')
    with pytest.raises(ValueError):register_round(path,1,'different')
    with pytest.raises(ValueError):register_round(path,4,'fourth')
    state=tmp_path/'evaluation.json';reserve_evaluation(state,'data','human')
    with pytest.raises(ValueError):reserve_evaluation(state,'data','human')


def test_failure_predictions_count_as_false_negative_with_wilson_n():
    from agent_ops.routing_metrics import classification_metrics
    report=classification_metrics(['order_status','order_status'],['order_status',None])
    assert report['accuracy']==.5 and report['prediction_failures']==1
    assert report['per_class']['order_status']['recall']==.5
    assert report['per_class']['order_status']['recall_wilson']['n']==2
    assert report['per_class']['order_status']['precision_wilson']['n']==1
    assert report['prediction_label_order'][-1]=='FAILED'


def test_freeze_tag_precedes_selection_and_rejects_source_changes(tmp_path,monkeypatch):
    import hashlib
    import json
    import subprocess

    from agent_ops.p8_release import verify_freeze
    monkeypatch.chdir(tmp_path)
    def git(*args):return subprocess.check_output(['git',*args],stderr=subprocess.DEVNULL)
    git('init');git('config','user.email','fixture@example.invalid');git('config','user.name','fixture')
    (tmp_path/'policy.py').write_text('frozen policy')
    (tmp_path/'freeze.json').write_text(json.dumps({'stage':'frozen_before_final_selection','files':{
        'policy.py':hashlib.sha256(b'frozen policy').hexdigest()}}))
    git('add','.');git('commit','-m','fixture');git('tag','-a','fixture-freeze','-m','fixture')
    assert verify_freeze('fixture-freeze',tmp_path.joinpath('freeze.json').relative_to(tmp_path))['tag']=='fixture-freeze'
    (tmp_path/'policy.py').write_text('changed policy')
    with pytest.raises(ValueError,match='SOURCE_CHANGED'):verify_freeze('fixture-freeze',tmp_path.joinpath('freeze.json').relative_to(tmp_path))


def test_final_workbook_blind_and_incomplete_annotations_block(tmp_path):
    import runpy

    from openpyxl import load_workbook
    prep=runpy.run_path('scripts/prepare-p8-final.py',run_name='fixture')
    evaluation=runpy.run_path('scripts/evaluate-p8-final.py',run_name='fixture')
    rows=[{'id':f'fixture-{i}','source':'ecd_redacted','message':'=literal text','context':[]} for i in range(200)]
    path=tmp_path/'blind.xlsx';prep['workbook'](path,rows)
    b=load_workbook(path);sheet=b['02_最终盲标200条']
    assert b.sheetnames[0]=='01_标注说明' and len(b.sheetnames)==2
    assert sheet.max_row==201 and sheet.max_column==9
    for row in sheet.iter_rows(min_row=2):
        assert row[4].data_type=='s' and all(c.value is None for c in row[5:])
    with pytest.raises(ValueError,match='INCOMPLETE'):
        evaluation['labels_from_workbook'](path,path,rows)
