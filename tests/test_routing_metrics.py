import json
import runpy

import pytest
from openpyxl import Workbook

from agent_ops.routing_metrics import classification_metrics, group_metrics, propose_clarify_subtype


def test_hand_calculated_kappa_per_class_and_fixed_macro():
    report=classification_metrics(['order_status','order_status','return_request','return_request'],
                                  ['order_status','order_status','order_status','return_request'])
    assert report['accuracy']==.75 and report['cohen_kappa']==.5
    assert report['per_class']['order_status']['precision']==pytest.approx(2/3)
    assert report['per_class']['order_status']['recall']==1
    assert report['per_class']['order_status']['f1']==pytest.approx(.8)
    assert report['per_class']['return_request']['f1']==pytest.approx(2/3)
    assert report['macro_f1_fixed_four']==pytest.approx((.8+2/3)/4)
    assert report['macro_f1_present_gold']==pytest.approx((.8+2/3)/2)
    assert report['per_class']['human']['recall'] is None
    assert report['confusion_matrix_rows_gold_cols_prediction']==[[2,0,0,0],[1,1,0,0],[0,0,0,0],[0,0,0,0]]


def test_abstentions_remain_in_coverage_and_majority_tie_is_explicit():
    rows=[{'id':str(i),'source':'ecd_redacted','gold':label,'rule':'clarify','model':'clarify','hybrid':'clarify'}
          for i,label in enumerate(['clarify','order_status',None])]
    result=group_metrics(rows)['combined']
    assert result['total_rows']==3 and result['labeled_rows']==2 and result['abstentions']==1
    assert result['abstention_ids']==['2'] and result['coverage']==pytest.approx(2/3)
    assert result['versions']['model']['n']==2
    assert result['majority_baseline']['tied_classes']==['order_status','clarify']
    assert result['majority_baseline']['selected_class_fixed_order']=='order_status'


def test_degenerate_kappa_and_unknown_labels_not_silently_changed():
    assert classification_metrics(['clarify'],['clarify'])['cohen_kappa'] is None
    assert classification_metrics([],[])['accuracy'] is None
    with pytest.raises(ValueError):classification_metrics(['not_a_label'],['clarify'])
    with pytest.raises(ValueError):classification_metrics(['clarify'],[])


def test_clarify_proposal_preserves_root_and_prioritizes_explicit_scope():
    assert propose_clarify_subtype('商品规格咨询，末尾致谢')[0]=='out_of_scope'
    assert propose_clarify_subtype('致谢收尾')[0]=='no_action'
    assert propose_clarify_subtype('否定之后没有明确需求')[0]=='insufficient_or_ambiguous'


def test_workbook_gold_import_validates_blind_content_and_preserves_abstention(tmp_path):
    module=runpy.run_path('scripts/analyze-p8-human-development.py',run_name='fixture_importer')
    original=Workbook();original.active.title='02_盲标100条';sheet=original.active
    sheet.append(['id','source','role','context','message','human role','gold','note'])
    for i in range(100):sheet.append([f'p2-{i}','fixture','unknown','fixture','fixture','','',''])
    before=tmp_path/'original.xlsx';original.save(before)
    for i in range(2,102):sheet.cell(i,6,'顾客');sheet.cell(i,7,'clarify')
    sheet.cell(2,6,'无法确认');sheet.cell(2,7).value=None
    annotated=tmp_path/'annotated.xlsx';original.save(annotated)
    dataset=tmp_path/'dataset.json';dataset.write_text(json.dumps({'samples':[{'id':f'p2-{i}','source':'p2_synthetic'} for i in range(100)]}))
    predictions=tmp_path/'predictions.json';predictions.write_text(json.dumps([{'id':f'p2-{i}','rule':'clarify','model':'clarify','hybrid':'clarify'} for i in range(100)]))
    rows=module['read_rows'](annotated,before,dataset,predictions)
    assert len(rows)==100 and rows[0]['gold'] is None and sum(r['gold'] is not None for r in rows)==99
    sheet.cell(2,5,'changed');original.save(annotated)
    with pytest.raises(ValueError,match='CONTENT_MODIFIED'):
        module['read_rows'](annotated,before,dataset,predictions)
