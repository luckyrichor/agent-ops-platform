"""Post-hoc offline inference and attribution. Never execute frozen routing or any API."""
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from agent_ops.p8_release import verify_freeze
from agent_ops.routing_metrics import wilson_interval
from agent_ops.routing_posthoc import VERSIONS, holm_adjust, paired_analysis

OUT=Path('docs/reports/p8-final-followup')


def attribution(row, note):
    text=row['message'];context=' '.join(row['context'])
    address=bool(re.search(r'地址|收件人|收货地',text) and re.search(r'改|换|修改',text))
    carrier=bool(re.search(r'快递|物流|承运|运费|包邮|顺丰|申通|圆通|中通|韵达|邮政|EMS|ems',text))
    dispatch=bool(re.search(r'发货|寄出|发走|发出|安排发',text))
    urging=bool(re.search(r'催|尽快|快点|赶紧|抓紧|加急|早[点些]|急用|着急|还没|未发|没发',text))
    note_urging=bool(re.search(r'催.*发|要求.*发货|催发货',note))
    time=bool(re.search(r'几时|何时|什么时候|啥时候|多久|几天|哪天|今天|明天|后天|到货|送到|能到|到达|收到|时间',text))
    strong_context=bool(re.search(r'已经付款|付款成功|已付款|已支付|订单编号|订单号|已发货|已经发出|快递单号|物流单号|包裹.*已',context))
    if address:category='address_change'
    elif carrier:category='carrier_related'
    elif note_urging or (dispatch and urging):category='urge_dispatch'
    elif row['subtype']=='无需处理':category='other'
    elif time and not row['context']:category='time_without_context'
    else:category='other'
    time_ambiguous=(category=='time_without_context' or (category=='other' and time and not strong_context and row['subtype']!='无需处理'))
    interpretation=('boundary_ambiguous_not_relabelled' if time_ambiguous else
      'clear_scope_or_no_action_error' if category in {'address_change','carrier_related','urge_dispatch'} or row['subtype']=='无需处理' else
      'insufficient_semantic_evidence')
    return {'category':category,'interpretation':interpretation,'actual_context_empty':not row['context'],
      'annotation_subtype':row['subtype'],'time_with_nondecisive_context':time_ambiguous and bool(row['context']),'strong_fulfillment_context_pattern':strong_context,
      'annotation_note_urging_pattern':note_urging}


def hybrid_reason(row):
    if row['hybrid_model_used']:
        t=row['message'].lower()
        if any(word in t for word in ['不','别','不要','并','同时','如果','ignore',"don't",'do not']):return 'model_branch_negation_condition_or_mixed_marker'
        return 'model_branch_no_unique_keyword_group'
    # Stored confidence 1 can arise from the carrier boundary override or one keyword group.
    from agent_ops.routing_policy import CARRIER_OPTIONS, PROGRESS
    if CARRIER_OPTIONS.search(row['message']) and not PROGRESS.search(row['message']):return 'rule_branch_carrier_option_override'
    return 'rule_branch_one_keyword_group'


def main():
    OUT.mkdir(parents=True,exist_ok=True)
    frozen=verify_freeze('p8-routing-freeze-20261030',Path('evals/p8/freeze.json'))
    root=Path('.local/p8-final-release');data_hash=hashlib.sha256((root/'final-dataset.json').read_bytes()).hexdigest()
    state=Path('.local/p8-final-evaluations')/(data_hash+'.json');before=state.read_bytes();ledger=json.loads(before)
    assert ledger['stage']=='completed'
    assert ledger['report_sha256']==hashlib.sha256(Path('docs/reports/p8-final/metrics.json').read_bytes()).hexdigest()
    rows=json.loads((root/'final-scored-private.json').read_text());valid=[r for r in rows if r['label'] is not None]
    report={'analysis':'posthoc_offline_only','frozen_commit':frozen['commit'],'final_evaluation_rerun':False,
      'bootstrap_replicates':20000,'seed':20261030,'primary_metrics':['agreement','macro_f1'],
      'new_exploratory_metrics':['human_recall','return_recall','false_human_rate','human_false_discovery_rate'],
      'comparison_policy':'Do not base version claims on the two ECD human-gold rows; use human-excluded sensitivity.',
      'bootstrap_limit':'Conditional on fixed observed gold counts; finite stratified sample, no population prevalence inference, no provider reruns.',
      'groups':{}}
    comparisons=[]
    for source in ['ecd_redacted','synthetic_rewrite']:
        group=[r for r in valid if r['source']==source]
        full=paired_analysis(group,replicates=20000,seed=20261030,exclude_gold_human=False)
        three=paired_analysis(group,replicates=20000,seed=20261031,exclude_gold_human=True)
        for result in (full,three):comparisons.extend(result['pairs'])
        exploratory={}
        for v in VERSIONS:
            human_tp=sum(r['label']==r[v]=='human' for r in group)
            human_n=sum(r['label']=='human' for r in group)
            predicted_human=sum(r[v]=='human' for r in group);false_human=predicted_human-human_tp
            return_tp=sum(r['label']==r[v]=='return_request' for r in group)
            return_n=sum(r['label']=='return_request' for r in group)
            exploratory[v]={'human_recall':wilson_interval(human_tp,human_n),
              'return_recall':wilson_interval(return_tp,return_n),
              'false_human_rate_nonhuman_denominator':wilson_interval(false_human,len(group)-human_n),
              'human_false_discovery_rate_predicted_human_denominator':wilson_interval(false_human,predicted_human)}
        branch=Counter('model' if r['hybrid_model_used'] else 'rule' for r in group)
        report['groups'][source]={'full_four_class':full,'exclude_gold_human_three_class':three,
          'exploratory_not_prespecified':exploratory,'hybrid':{'branches':dict(branch),
            'diff_from_rule':sum(r['hybrid']!=r['rule'] for r in group),'diff_from_model':sum(r['hybrid']!=r['model'] for r in group),
            'rule_model_disagreements':sum(r['rule']!=r['model'] for r in group)}}
    adjusted=holm_adjust([p['mcnemar_exact_two_sided_p'] for p in comparisons])
    for pair,p in zip(comparisons,adjusted,strict=True):pair['mcnemar_holm_p_12_comparisons']=p
    book=load_workbook(root/'annotated.xlsx',data_only=True)['02_最终盲标200条']
    notes={book.cell(i,1).value:book.cell(i,9).value or '' for i in range(2,202)}
    attribution_rows=[]
    for row in valid:
        if row['source']!='ecd_redacted' or row['label']!='clarify' or not any(row[v]=='order_status' for v in ['rule','model']):continue
        attribution_rows.append({'id':row['id'],'model_error':row['model']=='order_status','rule_error':row['rule']=='order_status',**attribution(row,notes[row['id']])})
    assert sum(r['model_error'] for r in attribution_rows)==29 and sum(r['rule_error'] for r in attribution_rows)==13
    report['attribution']={'method':'Local deterministic semantic-feature and original annotation-note assisted coding; no new human adjudication or gold changes.',
      'time_category_definition':'Strict time_without_context requires an empty context array. Time questions with present but nondecisive fulfillment context remain other, separately flagged as ambiguous.',
      'categories':{v:dict(Counter(r['category'] for r in attribution_rows if r[v+'_error'])) for v in ['model','rule']},
      'interpretations':{v:dict(Counter(r['interpretation'] for r in attribution_rows if r[v+'_error'])) for v in ['model','rule']},
      'strict_empty_context_count':{v:sum(r[v+'_error'] and r['actual_context_empty'] for r in attribution_rows) for v in ['model','rule']}}
    with (OUT/'clarify-order-attribution.csv').open('w') as f:
        w=csv.DictWriter(f,fieldnames=list(attribution_rows[0]));w.writeheader();w.writerows(sorted(attribution_rows,key=lambda r:r['id']))
    hybrid_rows=[]
    for r in valid:
        if r['rule']==r['model']==r['hybrid']:continue
        result='model_branch_inherits_model' if r['hybrid_model_used'] else 'rule_branch_corrects_model' if r['hybrid']==r['label'] else 'rule_branch_overrides_correct_model' if r['model']==r['label'] else 'rule_branch_both_wrong'
        hybrid_rows.append({'id':r['id'],'source':r['source'],'gold':r['label'],'rule':r['rule'],'model':r['model'],'hybrid':r['hybrid'],
            'stored_rule_confidence':r['rule_confidence'],'threshold':r['threshold'],'hybrid_model_used':r['hybrid_model_used'],
            'branch_reason':hybrid_reason(r),'outcome':result})
    with (OUT/'hybrid-disagreements.csv').open('w') as f:
        w=csv.DictWriter(f,fieldnames=list(hybrid_rows[0]));w.writeheader();w.writerows(sorted(hybrid_rows,key=lambda r:r['id']))
    report['hybrid_disagreement_rows']=len(hybrid_rows)
    assert state.read_bytes()==before
    (OUT/'analysis.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    print(json.dumps({'posthoc_completed':True,'model_calls':0,'final_state_unchanged':True,'attribution':report['attribution'],
      'groups':{s:{'hybrid':g['hybrid'],'three_macro':g['exclude_gold_human_three_class']['versions'],
        'three_pairs':g['exclude_gold_human_three_class']['pairs']} for s,g in report['groups'].items()}},ensure_ascii=False))


if __name__=='__main__':main()
