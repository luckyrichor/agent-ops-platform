"""The final evaluation is unavailable until the owner returns a fully annotated blind workbook."""
import argparse
import asyncio
import hashlib
import json
import os
import shlex
from datetime import UTC, datetime
from pathlib import Path

from openpyxl import load_workbook
from pydantic import ValidationError

from agent_ops.deepseek import DeepSeekConfig, DeepSeekEvaluation, EvaluationUnavailable
from agent_ops.p8_pilot import PilotAnswer, private_write
from agent_ops.p8_release import reserve_evaluation, verify_freeze
from agent_ops.routing_evaluation import DevelopmentSample, development_versions, request_messages
from agent_ops.routing_metrics import LABELS, classification_metrics


def labels_from_workbook(path,original,rows):
    a=load_workbook(path,data_only=False,keep_links=False)['02_最终盲标200条']
    b=load_workbook(original,data_only=False,keep_links=False)['02_最终盲标200条']
    if a.max_row!=201 or a.max_column!=9:raise ValueError('FINAL_EXPECTS_200_ROWS')
    gold={}
    for i,row in enumerate(rows,2):
        if [a.cell(i,j).value for j in range(1,6)]!=[b.cell(i,j).value for j in range(1,6)]:raise ValueError('FINAL_SOURCE_CHANGED')
        if any(a.cell(i,j).data_type=='f' for j in range(1,10)):raise ValueError('FORMULA_NOT_ALLOWED')
        role=a.cell(i,6).value;label=a.cell(i,7).value
        if role not in {'顾客','客服','无法确认'}:raise ValueError('ANNOTATION_INCOMPLETE')
        if role=='顾客' and label not in LABELS:raise ValueError('ANNOTATION_INCOMPLETE')
        if role!='顾客' and label is not None:raise ValueError('NONCUSTOMER_HAS_GOLD')
        subtype=a.cell(i,8).value
        if subtype not in {None,'信息不足','超范围咨询','无需处理'}:raise ValueError('UNKNOWN_SUBTYPE')
        gold[row['id']]={'label':label,'role':role,'subtype':subtype}
    return gold


async def evaluate(args):
    frozen=verify_freeze(args.freeze_tag,Path('evals/p8/freeze.json'))
    data_path=args.release_dir/'final-dataset.json';data=json.loads(data_path.read_text())
    if data['freeze_commit']!=frozen['commit']:raise ValueError('FINAL_FREEZE_MISMATCH')
    rows=data['rows'];gold=labels_from_workbook(args.annotated,args.release_dir/'blind-final.xlsx',rows)
    if len(rows)!=200:raise ValueError('FINAL_EXPECTS_200_ROWS')
    digest=hashlib.sha256(data_path.read_bytes()).hexdigest()
    preparation=json.loads((args.release_dir/'preparation-summary.json').read_text())
    if preparation['dataset_sha256']!=digest or preparation['workbook_sha256']!=hashlib.sha256((args.release_dir/'blind-final.xlsx').read_bytes()).hexdigest():
        raise ValueError('FINAL_RELEASE_CHANGED')
    state=Path('.local/p8-final-evaluations')/(digest+'.json')
    config=DeepSeekConfig.from_env()
    if any(getattr(config,key)!=frozen[key] for key in ['model','endpoint','concurrency','max_attempts','timeout_seconds']):
        raise ValueError('FINAL_PROVIDER_CONFIG_CHANGED')
    reserve_evaluation(state,digest,hashlib.sha256(args.annotated.read_bytes()).hexdigest())
    service=DeepSeekEvaluation(config,args.release_dir/'evaluation-cache',allow_network=True,allow_ecd=True)
    queue=asyncio.Queue()
    for row in rows:queue.put_nowait(row)
    scored=[]
    async def worker():
        while not queue.empty():
            row=queue.get_nowait();answer=gold[row['id']];model=None;failure=None
            if answer['label'] is not None:
                messages=request_messages(DevelopmentSample.model_validate(row),'classifier')
                private_write(args.release_dir/'evaluation-requests'/(row['id']+'.json'),messages)
                try:
                    result=await service.complete(dataset_digest=digest,source=row['source'],messages=messages)
                    private_write(args.release_dir/'evaluation-returns'/(row['id']+'.json'),result)
                    model=PilotAnswer.model_validate_json(result['content']).route
                except (EvaluationUnavailable,ValidationError) as e:failure=type(e).__name__
            versions=development_versions(DevelopmentSample.model_validate(row),model)
            scored.append({**row,**answer,**versions,'model_failure':failure})
            queue.task_done()
    try:
        await asyncio.gather(*(worker() for _ in range(config.concurrency)))
        private_write(args.release_dir/'final-scored-private.json',scored)
        report={'stage':'single_final_evaluation','freeze_tag':args.freeze_tag,'freeze_commit':frozen['commit'],
                'model':config.model,'endpoint':config.endpoint,'temperature':0,'reported_at':datetime.now(UTC).isoformat(),
                'samples':200,'network_requests':service.network_requests,'retries':service.retries,'rate_limits':service.rate_limits,
                'groups':{},'interpretation':'ECD and generated paraphrases reported separately; no natural-population combined claim',
                'wilson':'nominal 95% binomial Wilson with numerator/denominator; no weighted pseudo-count intervals'}
        for source in ['ecd_redacted','synthetic_rewrite']:
            items=[r for r in scored if r['source']==source];valid=[r for r in items if r['label'] is not None]
            report['groups'][source]={'rows':len(items),'labeled':len(valid),'abstentions':len(items)-len(valid),
                'versions':{v:classification_metrics([r['label'] for r in valid],[r[v] for r in valid]) for v in ['rule','model','hybrid']},
                'gold_minimum_30_each':{label:sum(r['label']==label for r in valid)>=30 for label in LABELS}}
        public=Path('evals/p8')
        with (public/'final-ecd-human.csv').open('w') as f:
            f.write('id,label\n')
            for row in sorted(scored,key=lambda r:r['id']):
                if row['source']=='ecd_redacted':f.write(row['id']+','+(row['label'] or '')+'\n')
        with (public/'final-synthetic.jsonl').open('w') as f:
            for row in sorted(scored,key=lambda r:r['id']):
                if row['source']=='synthetic_rewrite':f.write(json.dumps({k:row[k] for k in ['id','source','message','context','label','subtype']},ensure_ascii=False)+'\n')
        report_path=Path('docs/reports/p8-final/metrics.json');report_path.parent.mkdir(parents=True,exist_ok=True)
        report_path.write_text(json.dumps(report,ensure_ascii=False,indent=2))
        private_write(state,{'stage':'completed','dataset_sha256':digest,'annotation_sha256':hashlib.sha256(args.annotated.read_bytes()).hexdigest(),'report_sha256':hashlib.sha256(report_path.read_bytes()).hexdigest()})
        print(json.dumps({'final_evaluation_completed':True,'samples':200,'report':'docs/reports/p8-final/metrics.json'}))
    finally:await service.aclose()


def main():
    p=argparse.ArgumentParser(__doc__);p.add_argument('--freeze-tag',required=True);p.add_argument('--release-dir',type=Path,required=True)
    p.add_argument('--annotated',type=Path,required=True);p.add_argument('--key-file',type=Path,required=True)
    a=p.parse_args();os.umask(0o077)
    for path in [a.release_dir,a.annotated]:
        if path.is_symlink() or not path.resolve().is_relative_to((Path.cwd()/'.local').resolve()):p.error('private inputs required')
    if a.key_file.stat().st_mode&0o777!=0o600:p.error('key environment mode600 required')
    for line in a.key_file.read_text().splitlines():
        for word in shlex.split(line,comments=True):
            if word.startswith('DEEPSEEK_API_KEY='):os.environ['DEEPSEEK_API_KEY']=word.split('=',1)[1]
    asyncio.run(evaluate(a))


if __name__=='__main__':
    try:main()
    except Exception as e:  # noqa: BLE001 -- final text/labels and credentials remain private
        print(json.dumps({'evaluation_failed_type':type(e).__name__}));raise SystemExit(1) from None
