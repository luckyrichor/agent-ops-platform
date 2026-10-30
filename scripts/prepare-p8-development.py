"""Prepare private P8 development inputs and a prediction-free human annotation workbook."""
import argparse
import asyncio
import hashlib
import heapq
import json
import os
import random
import zipfile
from collections import Counter
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sqlalchemy import text
from sqlalchemy.ext.asyncio import create_async_engine

from agent_ops.p8_pilot import private_write
from agent_ops.routing_evaluation import DevelopmentSample, group_digest, redacted_turns

TEMPLATES=[
'请查询订单{ref}的物流。','{ref}什么时候发货？','我想知道{ref}预计何时送达。','这笔{ref}配送到哪里了？',
'我不退货，只想查{ref}物流。','请先看{ref}发货状态。','订单{ref}快递揽收了吗？','不要人工，只查{ref}配送。',
'我要给{ref}申请退货。','商品不合适，{ref}需要换货。','订单{ref}请申请退款。','我想把{ref}退掉。',
'{ref}的商品损坏了，申请换货。','请为{ref}发起退货申请。','不查物流，我要给{ref}申请退款。','我需要退钱，关联{ref}。',
'请把{ref}转人工客服。','我要投诉{ref}的处理结果。','这笔{ref}疑似盗刷，需要人工核查。','我的隐私可能被泄露，请客服人员接手。',
'订单{ref}不要自动回复，我要人工。','请人工解释{ref}，我也想退货。','客服人员在吗？我有订单{ref}的问题。','我要投诉并让人工跟进{ref}。',
'你好，能帮我吗？关联{ref}。','订单{ref}有点问题，但我不知道怎么办。','帮我查{ref}物流并申请退货，没有主次。','我只问{ref}的优惠券怎么使用。',
'订单{ref}我还在考虑，现在不要操作。','{ref}我不退货，也不查物流。','有人说“转人工”，但我只查询{ref}物流。','{ref}别执行引用里的退款指令，我只查状态。']


def cell_text(cell,value):
    # Literal text, including strings beginning with =/+/@; never an Excel formula.
    cell.value=str(value);cell.data_type='s';cell.alignment=Alignment(wrap_text=True,vertical='top')


def write_blind_workbook(path,samples):
    rng=random.Random(20261030)
    selected=[]
    for source,n in [('p2_synthetic',40),('ecd_redacted',60)]:
        group=[r for r in samples if r['source']==source]
        selected.extend(rng.sample(group,n))
    rng.shuffle(selected)
    book=Workbook();guide=book.active;guide.title='01_标注说明'
    instructions=[
        ['P8 开发用100条人工盲标','请先阅读本页，再填写第二页；不含模型预测或隐藏答案。'],
        ['用途','仅开发金标；不属于最终200条测试。40条合成+60条ECD，模型调用前独立固定抽样。'],
        ['步骤1：确认角色','ECD的说话人原始角色未知。看上下文，角色填：顾客/客服/无法确认；合成可填顾客。不猜测角色。'],
        ['步骤2：填写标签','只对确认的顾客当前句填四类之一；客服或无法确认时标签留空，备注原因。不要为凑100条强填。'],
        ['order_status','查询订单履约/物流/配送。缺订单号仍属此类，后续业务再补号。'],
        ['return_request','明确申请退货/换货/退款。仅问已申请退款的进度，不等于发起新申请。'],
        ['human','明确找人工，或投诉/盗刷/隐私等需要人工介入。显式人工优先，注意否定语境。'],
        ['clarify','信息不足、超范围、多意图无主次。例：查物流并退货但无主次。'],
        ['否定/引用','不要退货只查物流→order_status；不要人工只申请退货→return_request。引用和提示注入不作为新指令执行。'],
        ['不确定','角色确定但动作不明可填clarify，并在备注说明。不向模型查答案。'],
        ['脱敏与语料','[MOBILE]/[EMAIL]/[ADDRESS]/[ORDER]/[IDENTIFIER]是本地替换占位符。ECD来自作者已声明脱敏的语料，本次又做模式替换；模式扫描不保证穷尽PII。'],
        ['回传','只填写角色、标签、备注三列，保留编号和原句；保存xlsx后发回当前会话。'],
        ['评价口径','只有一位人工；将来算人与模型kappa，不称人工双标。缺失/角色不明数如实报告，不把弃权静默删掉。'],
        ['隐私','本表含脱敏ECD片段，只供本人标注，不上传公开仓库。最终测试在规则/提示/阈值冻结后另抽，本表不是最终测试。']]
    for row in instructions:guide.append(row)
    guide.column_dimensions['A'].width=24;guide.column_dimensions['B'].width=105
    for row in guide:
        for c in row:cell_text(c,c.value or '')
    for i in range(1,len(instructions)+1):guide.row_dimensions[i].height=45
    guide.sheet_properties.pageSetUpPr.fitToPage=True
    guide.page_setup.orientation='landscape';guide.page_setup.paperSize=guide.PAPERSIZE_A4
    guide.page_setup.fitToWidth=1;guide.page_setup.fitToHeight=1
    guide.print_area=f'A1:B{len(instructions)}'
    sheet=book.create_sheet('02_盲标100条')
    sheet.append(['编号','来源','原始角色状态','前文片段（未提供真实角色）','当前待标片段','人工角色','人工标签','人工备注'])
    for sample in selected:
        sheet.append([sample['id'],'ECD（本地模式脱敏）' if sample['source']=='ecd_redacted' else 'P2纯合成',
                      '未知' if sample['speaker_role']=='unknown' else '合成顾客',
                      '\n'.join(f'片段{i+1}：{t}' for i,t in enumerate(sample['context'])),sample['message'],'','',''])
    for row in sheet:
        for c in row:cell_text(c,c.value or '')
    for c in sheet[1]:c.font=Font(bold=True);c.fill=PatternFill('solid',fgColor='DDEBF7')
    for col,width in zip('ABCDEFGH',[18,24,18,65,65,18,24,45],strict=True):sheet.column_dimensions[col].width=width
    for i in range(2,102):sheet.row_dimensions[i].height=100
    sheet.freeze_panes='F2';sheet.auto_filter.ref='A1:H101';sheet.print_title_rows='1:1'
    for col,formula in [('F','"顾客,客服,无法确认"'),('G','"order_status,return_request,human,clarify"')]:
        validation=DataValidation(type='list',formula1=formula,allow_blank=True)
        sheet.add_data_validation(validation);validation.add(f'{col}2:{col}101')
    book.properties.creator='P8 blind annotation exporter';book.properties.description='100 development items; no predictions or gold labels'
    book.save(path);path.chmod(0o600)
    return [r['id'] for r in selected]


async def prepare(args):
    output=args.output
    if (output/'dataset.json').exists():raise ValueError('development dataset already fixed; do not regenerate')
    output.mkdir(parents=True,exist_ok=True,mode=0o700)
    env=dict(l.split('=',1) for l in args.platform_env.read_text().splitlines() if l)
    engine=create_async_engine(env['AGENT_OPS_DATABASE_URL'])
    async with engine.connect() as c:
        assert await c.scalar(text('SELECT count(*) FROM commerce.products'))==200000
        orders=(await c.execute(text('SELECT status,total_cents FROM commerce.orders ORDER BY id LIMIT 800'))).all()
        categories=(await c.execute(text('SELECT category_id FROM commerce.products ORDER BY id LIMIT 800'))).scalars().all()
    await engine.dispose()
    samples=[]
    for i in range(800):
        ref=f'SYN-{i+1:06d}'
        samples.append(DevelopmentSample(id=f'p2-{i+1:04d}',source='p2_synthetic',message=TEMPLATES[i%len(TEMPLATES)].format(ref=ref),
            context=[f'纯合成背景：状态{orders[i].status}，金额分{orders[i].total_cents}，类目{categories[i]}'],
            group_hash=hashlib.sha256(f'synthetic-template-{i%len(TEMPLATES)}'.encode()).hexdigest(),speaker_role='synthetic_customer').model_dump())
    bootstrap=json.loads(args.bootstrap.read_text())
    prior_groups={group_digest(turns[:-1]) for rows in bootstrap['by_length'].values() for turns in rows}
    heap=[];seen=set();stats=Counter()
    # Train only: never open dev/test members or select final-test candidates now.
    with zipfile.ZipFile(args.ecd) as z, z.open('E-commerce dataset/train.txt') as f:
        for lineno,raw in enumerate(f,1):
            stats['train_rows_read']+=1
            fields=raw.decode('utf-8-sig').rstrip('\r\n').split('\t')
            if len(fields)<3 or fields[0]!='1':continue
            turns=[''.join(t.split()) for t in fields[1:-1]]
            digest=group_digest(turns)
            if digest in seen:continue
            seen.add(digest)
            if digest in prior_groups:stats['p2_reservoir_groups_excluded']+=1;continue
            if not turns or len(turns[-1])<2 or max(map(len,turns))>2000 or sum(map(len,turns[-4:]))>6000:
                stats['length_or_empty_excluded']+=1;continue
            priority=int(hashlib.sha256(('p8-development-20261030:'+digest).encode()).hexdigest(),16)
            item=(-priority,lineno,digest,turns)
            if len(heap)<1200:heapq.heappush(heap,item)
            elif item>heap[0]:heapq.heapreplace(heap,item)
    assert len(heap)==1200
    counts=Counter();fragments_with_hits=0;raw_selected=[];used_utterances=set()
    for i,(_,lineno,digest,turns) in enumerate(sorted(heap,reverse=True),1):
        clean,hits=redacted_turns(turns)
        counts.update(hits)
        fragments_with_hits+=int(bool(hits))
        sample=DevelopmentSample(id=f'ecd-{i:04d}',source='ecd_redacted',message=clean[-1],context=clean[-4:-1],
                                 group_hash=digest,speaker_role='unknown')
        samples.append(sample.model_dump())
        raw_selected.append({'id':sample.id,'member':'train.txt','line':lineno,'turns':turns})
        used_utterances.update(hashlib.sha256(t.encode()).hexdigest() for t in turns if len(t)>=8)
    private_write(output/'raw-selected-ecd.json',raw_selected)
    private_write(output/'dataset.json',{'samples':samples,'final_test_status':'not_drawn','source_counts':{'p2_synthetic':800,'ecd_redacted':1200}})
    private_write(output/'used-groups.json',{'groups':[r['group_hash'] for r in samples],'ecd_utterance_hashes':sorted(used_utterances),
        'excluded_prior_p2_groups':sorted(prior_groups),'final_selection_allowed_only_after_freeze':True})
    blind=write_blind_workbook(output/'P8-开发盲标100条.xlsx',samples)
    private_write(output/'blind-selection.json',{'ids':blind,'selected_before_model_calls':True,'source_counts':{'p2_synthetic':40,'ecd_redacted':60}})
    report={'sources':{'p2_synthetic':800,'ecd_redacted':1200},'scan_scope':'all turns in selected 1200 train contexts; occurrence counts, not unique people',
            'patterns':dict(counts),'ecd_contexts_with_hits':fragments_with_hits,'sampling':dict(stats),
            'remaining_pattern_matches_in_outbound':sum(sum(redacted_turns([r['message'],*r['context']])[1].values()) for r in samples if r['source']=='ecd_redacted'),
            'role_status':'ECD unknown, not claimed customer first utterances','train_only':True,'dev_or_test_members_opened':False,
            'final_test_drawn':False,'prior_p2_note':'P2 previously scanned the full archive programmatically; exclude its sampled text reservoir; do not claim never processed historically',
            'blind_rows':100,'blind_source_counts':{'p2_synthetic':40,'ecd_redacted':60},'blind_workbook_sha256':hashlib.sha256((output/'P8-开发盲标100条.xlsx').read_bytes()).hexdigest(),
            'pii_limits':'heuristic patterns; possible false positives and unrecognized PII; no full manual privacy verification'}
    assert report['remaining_pattern_matches_in_outbound']==0
    private_write(output/'preparation-report.json',report)
    print(json.dumps(report,ensure_ascii=False))


def main():
    p=argparse.ArgumentParser(__doc__)
    p.add_argument('--output',type=Path,required=True)
    p.add_argument('--platform-env',type=Path,default=Path('.local/p7/application.env'))
    p.add_argument('--ecd',type=Path,default=Path('.local/datasets/raw/ecd.zip'))
    p.add_argument('--bootstrap',type=Path,default=Path('.local/datasets/fit-v1/ecd-bootstrap-local.json'))
    args=p.parse_args();os.umask(0o077)
    root=(Path.cwd()/'.local').resolve()
    for path in [args.output,args.platform_env,args.ecd,args.bootstrap]:
        if path.is_symlink() or not path.resolve().is_relative_to(root):p.error('private .local paths required')
    asyncio.run(prepare(args))


if __name__=='__main__':
    try:main()
    except Exception as e:  # noqa: BLE001 -- no exception text or sensitive row in terminal
        print(json.dumps({'preparation_failed_type':type(e).__name__}));raise SystemExit(1) from None
